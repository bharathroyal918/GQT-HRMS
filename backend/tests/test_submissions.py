"""Comprehensive Automated Test Suite for Phase 7: Candidate Submission & Profile Dispatch (Steps 24-28).

Authoritative Workflow Coverage:
- Step 24: Prepare Candidate List (Job eligibility & Step 23 Placement Readiness Certification Gate)
- Step 25: Generate Candidate Roster Spreadsheet (Master Excel file with SHA-256 integrity digest)
- Step 26: Attach Resumes (Verify exact certified resume PDF versions and hashes)
- Step 27: Send Candidate Profiles to HR (Dispatch with transmission log & provider tracking)
- Step 28: Company Acknowledgement (Log recruiter receipt, feedback, drive date, callback timeline)
- Security: Role-Based Access Controls, Candidate Privacy & Data Isolation
- Audit: Immutable event logging for consequential actions
"""
import io
import hashlib
import uuid
from decimal import Decimal
import openpyxl
import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import Role, RoleCatalog, User, UserRoleAssignment
from apps.core.models import AuditLog
from apps.crm.models import Company, HRContact, CompanyDomainCategory
from apps.jobs.models import JobRequirement, JobRequirementStatus
from apps.students.models import (
    StudentProfile,
    StudentApplication,
    StudentApplicationStatus,
    StudentResume,
)
from apps.preparation.models import (
    PlacementReadinessCertification,
    PlacementReadinessStatus,
)
from apps.submissions.models import (
    SubmissionPackage,
    SubmissionPackageStatus,
    CandidateSubmission,
    CandidateSubmissionStatus,
    SubmissionDispatchLog,
    DispatchChannel,
    AcknowledgementChannel,
)
from apps.submissions.selectors import (
    get_eligible_candidates_for_job,
    get_submission_packages,
    get_submission_package_by_id,
    get_student_submissions,
)
from apps.submissions.services import (
    create_submission_package,
    generate_candidate_roster_spreadsheet,
    verify_and_attach_resumes,
    dispatch_candidate_profiles_to_hr,
    record_company_acknowledgement,
)


@pytest.fixture
def api_client():
    return APIClient()


@pytest.fixture(autouse=True)
def roles(db):
    roles_dict = {}
    for code, name in RoleCatalog.CHOICES:
        role, _ = Role.objects.get_or_create(code=code, defaults={"name": name, "is_active": True})
        roles_dict[code] = role
    return roles_dict


def make_user(email: str, role_code: str, roles) -> User:
    user = User.objects.create_user(
        email=email.lower().strip(),
        password="TestPassword123!",
        first_name="Test",
        last_name=role_code,
    )
    UserRoleAssignment.objects.create(
        user=user,
        role=roles[role_code],
        is_active=True,
    )
    return user


@pytest.fixture
def admin_user(db, roles):
    return make_user("admin.phase7@gqt.internal", RoleCatalog.ADMIN, roles)


@pytest.fixture
def placement_head(db, roles):
    return make_user("head.phase7@gqt.internal", RoleCatalog.PLACEMENT_HEAD, roles)


@pytest.fixture
def placement_officer(db, roles):
    return make_user("officer.phase7@gqt.internal", RoleCatalog.PLACEMENT_OFFICER, roles)


@pytest.fixture
def hr_relationship_mgr(db, roles):
    return make_user("hrmgr.phase7@gqt.internal", RoleCatalog.HR_RELATIONSHIP_MGR, roles)


@pytest.fixture
def student_user_a(db, roles):
    user = make_user("student.a.p7@gqt.edu", RoleCatalog.STUDENT, roles)
    user.first_name = "Candidate"
    user.last_name = "Alpha"
    user.save()
    return user


@pytest.fixture
def student_user_b(db, roles):
    user = make_user("student.b.p7@gqt.edu", RoleCatalog.STUDENT, roles)
    user.first_name = "Candidate"
    user.last_name = "Beta"
    user.save()
    return user


@pytest.fixture
def student_profile_a(db, student_user_a):
    return StudentProfile.objects.create(
        user=student_user_a,
        roll_number="1GQ21CS101",
        institutional_email=student_user_a.email,
        college_name="GQT Engineering College",
        degree="B.E.",
        department="Computer Science",
        graduation_year=2025,
        cgpa=Decimal("8.50"),
        active_backlogs=0,
        phone_number="+919876543210",
    )


@pytest.fixture
def student_profile_b(db, student_user_b):
    return StudentProfile.objects.create(
        user=student_user_b,
        roll_number="1GQ21CS102",
        institutional_email=student_user_b.email,
        college_name="GQT Engineering College",
        degree="B.E.",
        department="Information Science",
        graduation_year=2025,
        cgpa=Decimal("7.80"),
        active_backlogs=0,
        phone_number="+919876543211",
    )


def create_dummy_pdf_resume(student: StudentProfile, version: int = 1) -> StudentResume:
    content = b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF\n"
    digest = hashlib.sha256(content).hexdigest()
    uploaded = SimpleUploadedFile(f"resume_v{version}.pdf", content, content_type="application/pdf")
    return StudentResume.objects.create(
        student=student,
        version_number=version,
        file=uploaded,
        file_name=f"resume_v{version}.pdf",
        file_size_bytes=len(content),
        sha256_hash=digest,
        is_current=True,
    )


@pytest.fixture
def company(db):
    return Company.objects.create(
        name="Infosys Technologies",
        domain_category=CompanyDomainCategory.IT_SERVICES,
        city="Bengaluru",
        state="Karnataka",
    )


@pytest.fixture
def hr_contact(db, company):
    return HRContact.objects.create(
        company=company,
        name="Pooja Sharma",
        designation="Campus Talent Acquisition Lead",
        email="pooja.sharma@infosys.example.com",
        phone="+919888877777",
    )


@pytest.fixture
def approved_job(db, company, hr_contact):
    return JobRequirement.objects.create(
        company=company,
        hr_contact=hr_contact,
        title="Associate Software Engineer",
        job_description="Hiring fresher software engineers for IT delivery centers.",
        skills_required="Java, Python, SQL",
        status=JobRequirementStatus.APPROVED,
        min_salary=Decimal("450000.00"),
        max_salary=Decimal("650000.00"),
        registration_deadline=timezone.now() + timezone.timedelta(days=14),
    )


@pytest.fixture
def application_a(db, student_profile_a, approved_job):
    return StudentApplication.objects.create(
        student=student_profile_a,
        job_requirement=approved_job,
        status=StudentApplicationStatus.ELIGIBLE,
    )


@pytest.fixture
def application_b(db, student_profile_b, approved_job):
    return StudentApplication.objects.create(
        student=student_profile_b,
        job_requirement=approved_job,
        status=StudentApplicationStatus.ELIGIBLE,
    )


@pytest.fixture
def certified_resume_a(db, student_profile_a, placement_head):
    resume = create_dummy_pdf_resume(student_profile_a, version=1)
    PlacementReadinessCertification.objects.create(
        student=student_profile_a,
        approved_resume=resume,
        is_approved_for_placements=True,
        status=PlacementReadinessStatus.CERTIFIED,
        certified_by=placement_head,
        readiness_summary={"resume_score": 9, "mock_score": 36},
        notes="Candidate certified ready for IT software drives.",
    )
    return resume


# ==============================================================================
# Step 24: Candidate List & Package Creation Tests
# ==============================================================================

@pytest.mark.django_db
def test_01_eligible_candidates_selector_shows_readiness_certification_status(
    approved_job, student_profile_a, application_a, certified_resume_a, student_profile_b, application_b
):
    """Step 24 Selector: identifies eligible applicants and checks Step 23 Placement Head certification."""
    result = get_eligible_candidates_for_job(str(approved_job.id))
    assert result["job_id"] == str(approved_job.id)
    assert result["total_eligible"] == 2
    assert result["total_certified"] == 1  # Only student A has Step 23 certification

    candidates = result["candidates"]
    candidate_a = next(c for c in candidates if c["student_id"] == str(student_profile_a.id))
    assert candidate_a["has_readiness_certification"] is True
    assert candidate_a["approved_resume_version"] == 1
    assert candidate_a["can_be_submitted"] is True

    candidate_b = next(c for c in candidates if c["student_id"] == str(student_profile_b.id))
    assert candidate_b["has_readiness_certification"] is False
    assert candidate_b["can_be_submitted"] is False


@pytest.mark.django_db
def test_02_create_submission_package_success_step_24(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 24 Service: creates DRAFT submission package with certified candidate and frozen snapshot."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
        title="Infosys 2026 Batch - Cohort 1",
        notes="High-readiness candidates for campus drive.",
    )

    assert package.status == SubmissionPackageStatus.DRAFT
    assert package.candidate_count == 1
    assert package.company == approved_job.company
    assert package.hr_contact == approved_job.hr_contact

    # Verify CandidateSubmission record
    cs = package.candidates.first()
    assert cs.student == student_profile_a
    assert cs.submitted_resume == certified_resume_a
    assert cs.resume_version_number == 1
    assert cs.resume_sha256 == certified_resume_a.sha256_hash
    assert cs.status == CandidateSubmissionStatus.PENDING

    # Verify frozen snapshot
    snap = cs.candidate_snapshot
    assert snap["roll_number"] == "1GQ21CS101"
    assert snap["department"] == "Computer Science"
    assert snap["resume_version"] == 1

    # Verify AuditLog
    assert AuditLog.objects.filter(
        user=placement_officer,
        action="SUBMISSION_PACKAGE_CREATED",
        resource_id=str(package.id),
    ).exists()


@pytest.mark.django_db
def test_03_create_package_fails_if_candidate_lacks_readiness_certification_step_23_gate(
    approved_job, student_profile_b, application_b, placement_officer
):
    """Step 23 Gate Enforcement: rejects package creation if any candidate lacks Placement Head certification."""
    with pytest.raises(ValidationError) as exc:
        create_submission_package(
            job_requirement_id=str(approved_job.id),
            candidate_student_ids=[str(student_profile_b.id)],
            user=placement_officer,
        )
    assert "lacks Placement Head readiness certification (Workflow Step 23 Gate)" in str(exc.value)


@pytest.mark.django_db
def test_04_create_package_fails_if_job_not_approved(
    company, hr_contact, student_profile_a, placement_officer
):
    """Business Invariant: cannot submit candidates to a job that is still in DRAFT or unapproved."""
    draft_job = JobRequirement.objects.create(
        company=company,
        hr_contact=hr_contact,
        title="Draft Role",
        job_description="Unapproved description",
        skills_required="Python",
        status=JobRequirementStatus.DRAFT,
        registration_deadline=timezone.now() + timezone.timedelta(days=7),
    )
    with pytest.raises(ValidationError) as exc:
        create_submission_package(
            job_requirement_id=str(draft_job.id),
            candidate_student_ids=[str(student_profile_a.id)],
            user=placement_officer,
        )
    assert "Cannot create candidate submission for job in 'Draft" in str(exc.value)


@pytest.mark.django_db
def test_05_create_package_fails_with_empty_candidates(approved_job, placement_officer):
    """Validation: empty candidate list is rejected."""
    with pytest.raises(ValidationError) as exc:
        create_submission_package(
            job_requirement_id=str(approved_job.id),
            candidate_student_ids=[],
            user=placement_officer,
        )
    assert "At least one eligible candidate must be selected" in str(exc.value)


@pytest.mark.django_db
def test_06_create_package_prevents_duplicate_active_candidate_submission(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Deduplication / Conflict Protection: prevents adding candidate already in an active package for this job."""
    # Create first package
    create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )

    # Attempt second package with same student
    with pytest.raises(ValidationError) as exc:
        create_submission_package(
            job_requirement_id=str(approved_job.id),
            candidate_student_ids=[str(student_profile_a.id)],
            user=placement_officer,
        )
    assert "already included in an active submission package" in str(exc.value)


# ==============================================================================
# Step 25: Master Excel Spreadsheet Generation Tests
# ==============================================================================

@pytest.mark.django_db
def test_07_generate_roster_spreadsheet_success_step_25(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 25: generates master XLSX spreadsheet with student metrics & openpyxl styles."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )

    package = generate_candidate_roster_spreadsheet(str(package.id), user=placement_officer)

    assert package.roster_file is not None
    assert package.roster_file_name.endswith(".xlsx")
    assert len(package.package_hash) == 64  # SHA-256 digest
    assert package.roster_generated_at is not None

    # Inspect generated Excel content
    content = package.roster_file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws.title == "Candidate Roster"

    # Verify title and headers
    assert "GQT HRMS" in ws["A1"].value
    assert ws["A5"].value == "Sl. No."
    assert ws["B5"].value == "Roll Number"
    assert ws["C5"].value == "Candidate Name"

    # Verify student data row
    assert ws["B6"].value == "1GQ21CS101"
    assert ws["D6"].value == student_profile_a.institutional_email
    assert ws["G6"].value == "Computer Science"
    assert ws["K6"].value == "Resume v1"

    # Verify AuditLog
    assert AuditLog.objects.filter(
        user=placement_officer,
        action="ROSTER_SPREADSHEET_GENERATED",
        resource_id=str(package.id),
    ).exists()


@pytest.mark.django_db
def test_08_generate_roster_calculates_sha256_digest_and_stores_hash(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 25 Integrity: verifies package_hash exactly matches SHA-256 of the generated Excel bytes."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    package = generate_candidate_roster_spreadsheet(str(package.id), user=placement_officer)

    file_bytes = package.roster_file.read()
    computed_digest = hashlib.sha256(file_bytes).hexdigest()
    assert package.package_hash == computed_digest


# ==============================================================================
# Step 26: Resume Attachments Verification Tests
# ==============================================================================

@pytest.mark.django_db
def test_10_verify_and_attach_resumes_success_step_26(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 26: verifies certified resume files and advances package to READY_FOR_DISPATCH."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )

    package = verify_and_attach_resumes(str(package.id), user=placement_officer)
    assert package.status == SubmissionPackageStatus.READY_FOR_DISPATCH
    assert package.resumes_attached_at is not None

    assert AuditLog.objects.filter(
        user=placement_officer,
        action="SUBMISSION_RESUMES_ATTACHED",
        resource_id=str(package.id),
    ).exists()


@pytest.mark.django_db
def test_11_verify_and_attach_resumes_fails_on_hash_mismatch(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 26 Tamper Protection: rejects resume attachment if binary hash does not match certified snapshot."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )

    # Artificially alter resume sha256 to simulate tampered document
    certified_resume_a.sha256_hash = "0000000000000000000000000000000000000000000000000000000000000000"
    certified_resume_a.save()

    with pytest.raises(ValidationError) as exc:
        verify_and_attach_resumes(str(package.id), user=placement_officer)
    assert "Integrity check failed: resume hash" in str(exc.value)


# ==============================================================================
# Step 27: Profile Dispatch to Corporate HR Tests
# ==============================================================================

@pytest.mark.django_db
def test_12_dispatch_candidate_profiles_to_hr_success_step_27(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Step 27: dispatches package to corporate HR and transitions to DISPATCHED."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    generate_candidate_roster_spreadsheet(str(package.id), user=placement_officer)
    verify_and_attach_resumes(str(package.id), user=placement_officer)

    package = dispatch_candidate_profiles_to_hr(
        package_id=str(package.id),
        user=placement_officer,
        channel=DispatchChannel.EMAIL,
        recipient_email="pooja.sharma@infosys.example.com",
        recipient_name="Pooja Sharma",
        dispatch_notes="Batch 2026 profiles for upcoming drive.",
    )

    assert package.status == SubmissionPackageStatus.DISPATCHED
    assert package.dispatched_at is not None
    assert package.dispatched_by == placement_officer
    assert package.dispatch_channel == DispatchChannel.EMAIL
    assert package.recipient_email == "pooja.sharma@infosys.example.com"
    assert package.provider_reference.startswith("MSG-DISP-")

    # Check candidate submissions advanced to SUBMITTED
    cs = package.candidates.first()
    assert cs.status == CandidateSubmissionStatus.SUBMITTED

    # Check dispatch log record
    assert package.dispatch_logs.count() == 1
    log = package.dispatch_logs.first()
    assert log.status == "SENT"
    assert log.recipient == "pooja.sharma@infosys.example.com"
    assert log.provider_message_id == package.provider_reference

    # Verify AuditLog
    assert AuditLog.objects.filter(
        user=placement_officer,
        action="CANDIDATE_PROFILES_DISPATCHED",
        resource_id=str(package.id),
    ).exists()


@pytest.mark.django_db
def test_14_dispatch_fails_if_already_dispatched(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Idempotency / Concurrency Guard: prevents re-dispatching an already dispatched package."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    verify_and_attach_resumes(str(package.id), user=placement_officer)
    dispatch_candidate_profiles_to_hr(str(package.id), user=placement_officer)

    with pytest.raises(ValidationError) as exc:
        dispatch_candidate_profiles_to_hr(str(package.id), user=placement_officer)
    assert "Package has already been dispatched to corporate HR" in str(exc.value)


@pytest.mark.django_db
def test_15_dispatch_fails_if_required_recipient_contact_missing(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """Validation: email dispatch requires recipient email address."""
    # Blank out HR contact email
    approved_job.hr_contact.email = ""
    approved_job.hr_contact.save()

    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    with pytest.raises(ValidationError) as exc:
        dispatch_candidate_profiles_to_hr(
            str(package.id),
            user=placement_officer,
            channel=DispatchChannel.EMAIL,
            recipient_email="",
        )
    assert "recipient email address is required" in str(exc.value)


# ==============================================================================
# Step 28: Company Acknowledgement Tests
# ==============================================================================

@pytest.mark.django_db
def test_16_record_company_acknowledgement_success_step_28(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer, hr_relationship_mgr
):
    """Step 28: logs recruiter acknowledgment, drive date, and transitions to ACKNOWLEDGED."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    verify_and_attach_resumes(str(package.id), user=placement_officer)
    dispatch_candidate_profiles_to_hr(str(package.id), user=placement_officer)

    drive_date = timezone.now().date() + timezone.timedelta(days=14)
    callback_date = timezone.now().date() + timezone.timedelta(days=5)

    package = record_company_acknowledgement(
        package_id=str(package.id),
        user=hr_relationship_mgr,
        acknowledged_by_contact_name="Pooja Sharma",
        acknowledgement_channel=AcknowledgementChannel.EMAIL,
        recruiter_feedback="Profiles received. Candidate pool matches criteria. Shortlist in 5 days.",
        tentative_interview_date=drive_date,
        follow_up_callback_date=callback_date,
    )

    assert package.status == SubmissionPackageStatus.ACKNOWLEDGED
    assert package.acknowledged_at is not None
    assert package.acknowledged_by == hr_relationship_mgr
    assert package.acknowledged_by_contact_name == "Pooja Sharma"
    assert package.recruiter_feedback.startswith("Profiles received")
    assert package.tentative_interview_date == drive_date
    assert package.follow_up_callback_date == callback_date

    # Candidate submissions advanced to ACKNOWLEDGED
    cs = package.candidates.first()
    assert cs.status == CandidateSubmissionStatus.ACKNOWLEDGED

    # Verify AuditLog
    assert AuditLog.objects.filter(
        user=hr_relationship_mgr,
        action="COMPANY_ACKNOWLEDGEMENT_RECORDED",
        resource_id=str(package.id),
    ).exists()


@pytest.mark.django_db
def test_17_record_acknowledgement_fails_on_undispatched_package(
    approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """State Machine: cannot acknowledge a package that has not yet been dispatched."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    with pytest.raises(ValidationError) as exc:
        record_company_acknowledgement(
            str(package.id),
            user=placement_officer,
            acknowledged_by_contact_name="Recruiter Name",
        )
    assert "Cannot record company acknowledgment for package in 'Draft" in str(exc.value)


# ==============================================================================
# API Endpoints & Role-Based Access Control Tests
# ==============================================================================

@pytest.mark.django_db
def test_18_api_job_candidates_endpoint_permissions(
    api_client, approved_job, student_profile_a, application_a, certified_resume_a, placement_officer, student_user_a
):
    """API: Job candidates preview is accessible to Placement Staff, denied to students."""
    url = f"/api/v1/submissions/jobs/{approved_job.id}/candidates/"

    # Student denied
    api_client.force_authenticate(user=student_user_a)
    res = api_client.get(url)
    assert res.status_code == 403

    # Placement officer allowed
    api_client.force_authenticate(user=placement_officer)
    res = api_client.get(url)
    assert res.status_code == 200
    assert res.data["total_eligible"] == 1
    assert res.data["total_certified"] == 1
    assert len(res.data["candidates"]) == 1


@pytest.mark.django_db
def test_19_api_package_create_and_detail_lifecycle(
    api_client, approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """API: Complete submission package lifecycle via REST endpoints."""
    api_client.force_authenticate(user=placement_officer)

    # 1. Create Package (POST /api/v1/submissions/packages/)
    payload = {
        "job_requirement_id": str(approved_job.id),
        "candidate_student_ids": [str(student_profile_a.id)],
        "title": "API Test Package",
        "notes": "Testing REST API lifecycle.",
    }
    res = api_client.post("/api/v1/submissions/packages/", payload, format="json")
    assert res.status_code == 201
    package_id = res.data["id"]
    assert res.data["status"] == "DRAFT"
    assert res.data["candidate_count"] == 1

    # 2. Generate Roster (POST .../roster/generate/)
    res = api_client.post(f"/api/v1/submissions/packages/{package_id}/roster/generate/")
    assert res.status_code == 200
    assert len(res.data["package_hash"]) == 64

    # 3. Attach Resumes (POST .../resumes/attach/)
    res = api_client.post(f"/api/v1/submissions/packages/{package_id}/resumes/attach/")
    assert res.status_code == 200
    assert res.data["status"] == "READY_FOR_DISPATCH"

    # 4. Dispatch (POST .../dispatch/)
    dispatch_payload = {
        "channel": "EMAIL",
        "recipient_email": "hr@infosys.example.com",
        "recipient_name": "Pooja Sharma",
        "dispatch_notes": "Profiles sent via API test.",
    }
    res = api_client.post(f"/api/v1/submissions/packages/{package_id}/dispatch/", dispatch_payload, format="json")
    assert res.status_code == 200
    assert res.data["status"] == "DISPATCHED"
    assert res.data["provider_reference"].startswith("MSG-DISP-")

    # 5. Acknowledge (POST .../acknowledge/)
    ack_payload = {
        "acknowledged_by_contact_name": "Pooja Sharma",
        "acknowledgement_channel": "EMAIL",
        "recruiter_feedback": "Received successfully.",
    }
    res = api_client.post(f"/api/v1/submissions/packages/{package_id}/acknowledge/", ack_payload, format="json")
    assert res.status_code == 200
    assert res.data["status"] == "ACKNOWLEDGED"
    assert res.data["acknowledged_by_contact_name"] == "Pooja Sharma"


@pytest.mark.django_db
def test_20_api_roster_download_endpoint(
    api_client, approved_job, student_profile_a, application_a, certified_resume_a, placement_officer
):
    """API: Download generated candidate roster spreadsheet."""
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )
    generate_candidate_roster_spreadsheet(str(package.id), user=placement_officer)

    api_client.force_authenticate(user=placement_officer)
    res = api_client.get(f"/api/v1/submissions/packages/{package.id}/roster/download/")
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res["Content-Type"]
    assert "attachment;" in res["Content-Disposition"]


@pytest.mark.django_db
def test_22_api_student_submissions_data_isolation(
    api_client, approved_job, student_profile_a, application_a, certified_resume_a,
    student_profile_b, application_b, placement_officer, student_user_a, student_user_b
):
    """Candidate Isolation: Students can only view their own submissions on GET /my-submissions/."""
    # Package containing student A
    package = create_submission_package(
        job_requirement_id=str(approved_job.id),
        candidate_student_ids=[str(student_profile_a.id)],
        user=placement_officer,
    )

    # Student A queries my-submissions
    api_client.force_authenticate(user=student_user_a)
    res_a = api_client.get("/api/v1/submissions/my-submissions/")
    assert res_a.status_code == 200
    assert len(res_a.data) == 1
    assert res_a.data[0]["job_title"] == approved_job.title
    assert res_a.data[0]["company_name"] == approved_job.company.name

    # Student B queries my-submissions: should see 0
    api_client.force_authenticate(user=student_user_b)
    res_b = api_client.get("/api/v1/submissions/my-submissions/")
    assert res_b.status_code == 200
    assert len(res_b.data) == 0


@pytest.mark.django_db
def test_23_api_unauthenticated_requests_denied(api_client, approved_job):
    """Security: Unauthenticated requests are rejected with 401 Unauthorized."""
    res = api_client.get("/api/v1/submissions/packages/")
    assert res.status_code == 401

    res = api_client.get(f"/api/v1/submissions/jobs/{approved_job.id}/candidates/")
    assert res.status_code == 401

    res = api_client.get("/api/v1/submissions/my-submissions/")
    assert res.status_code == 401
